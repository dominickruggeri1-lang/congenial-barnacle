"""Leak sheet: recalculated values match the UE Systems curves and the report reader.

    python -m pytest tests -q
"""
import importlib.util
import os
import subprocess
import sys

import pytest
from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
from ald.leak_sheet import CURVES, build  # noqa: E402

REPORT = ("/root/.claude/skills/synced/bab096fc-fa00-4781-b9d9-622996c5b77d_86e901db-c2cf-4b96-9851-059f0f18a3c1"
          "/leak-detection-report/scripts/build_report.py")
RATE = 0.18
HRS = 24 * 365
KWH_PER_KCF = 3.29

# gas, psi, dB, repaired
LEAKS = [
    ("Air", 100, 30, "Y"),
    ("Air", 125, 42.5, "N"),     # fractional dB
    ("Air", 90, 25, "Y"),        # between 75 and 100 psi curves
    ("Nitrogen", 50, 20, "N"),
    ("Open Blow", 100, 60, "N"),
    ("Air", 150, 10, ""),        # blank repaired = not repaired
]


def curve(psi, db):
    pts = {p: a * db ** b for p, a, b in CURVES}
    if psi in pts:
        return pts[psi]
    lo = max(p for p in pts if p <= psi)
    hi = min(p for p in pts if p >= psi)
    return pts[lo] + (pts[hi] - pts[lo]) * (psi - lo) / (hi - lo)


@pytest.fixture(scope="module")
def sheet(tmp_path_factory):
    path = str(tmp_path_factory.mktemp("ls") / "Test Leak Sheet.xlsx")
    wb = build("Test Plant", rate=RATE, year=2026)
    ws = wb["Leak Log"]
    for i, (gas, psi, db, rep) in enumerate(LEAKS):
        r = 8 + i
        ws[f"B{r}"], ws[f"C{r}"], ws[f"D{r}"], ws[f"M{r}"] = gas, psi, db, rep or None
        ws[f"E{r}"], ws[f"I{r}"] = f"Loc {i+1}", f"Desc {i+1}"
    wb.save(path)
    out = subprocess.run([sys.executable, os.path.join(ROOT, "tools", "recalc.py"), path, "90"],
                         capture_output=True, text=True)
    assert '"status": "success"' in out.stdout, out.stdout + out.stderr
    return path


def test_cfm_cost_kwh(sheet):
    ws = load_workbook(sheet, data_only=True)["Leak Log"]
    air_rate = KWH_PER_KCF * RATE
    for i, (gas, psi, db, _) in enumerate(LEAKS):
        r = 8 + i
        cfm = curve(psi, db)
        assert ws[f"O{r}"].value == pytest.approx(cfm, rel=1e-9)
        per_kcf = 3.80 if gas == "Nitrogen" else air_rate
        assert ws[f"N{r}"].value == pytest.approx(cfm * 60 * HRS * per_kcf / 1000, rel=1e-9)
        kwh = cfm * 60 * HRS / 1000 * KWH_PER_KCF if gas in ("Air", "Open Blow") else 0
        assert ws[f"P{r}"].value == pytest.approx(kwh, rel=1e-9)
    assert ws["U10"].value.startswith("INTERPOLATED")
    assert ws["U8"].value == "Table curve"


def test_table_value_matches_legacy_chart():
    # Legacy Flow Rate Table displays 1.17 CFM at 18 dB, 150 psi and 0.80 at 18 dB, 100 psi.
    assert round(curve(150, 18), 2) == 1.17
    assert round(curve(100, 18), 2) == 0.80


def test_totals_and_percent_repaired(sheet):
    ws = load_workbook(sheet, data_only=True)["Leak Log"]
    air = [curve(p, d) for g, p, d, _ in LEAKS if g == "Air"]
    assert ws["A6"].value == pytest.approx(sum(air), rel=1e-9)
    # 5 leaks excluding the open blow, 2 marked Y. Old COUNTA formula would give 4/6.
    assert ws["O6"].value == pytest.approx(2 / 5)
    assert ws["S18"].value == pytest.approx(curve(100, 60), rel=1e-9)
    assert ws["S22"].value == 0


def test_report_skill_reads_it(sheet):
    if not os.path.exists(REPORT):
        pytest.skip("leak-detection-report skill not installed")
    spec = importlib.util.spec_from_file_location("build_report", REPORT)
    mod = importlib.util.module_from_spec(spec)
    sys.argv = ["x"]
    try:
        spec.loader.exec_module(mod)
    except SystemExit:
        pass
    leaks, blows, kwh_rate, air_rate = mod.load_sheet(sheet)
    assert kwh_rate == RATE
    assert len(blows) == 1 and len(leaks) == 5
    ws = load_workbook(sheet, data_only=True)["Leak Log"]
    assert sum(r["cfm"] for r in leaks if r["gas"] == "Air") == pytest.approx(ws["A6"].value)

"""Cost Engine: the generated Kawasaki job reproduces the hand-built ALD-2026-1013-KM workbook."""
import os
import subprocess
import sys

import pytest
from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
from ald.cost_engine import build, load  # noqa: E402

# Values read from "ALD Quote - Kawasaki Maryville Leak Detection and Repair - Cost Engine.xlsx"
EXPECTED = dict(survey=30750, repair=81250, total=112000, leak_cfm=1961, guaranteed_cfm=588,
                potential=270387.3864, guaranteed_savings=81074.8512)


@pytest.fixture(scope="module")
def engine(tmp_path_factory):
    job, rates = load(os.path.join(ROOT, "jobs", "kawasaki_maryville.yaml"))
    wb, cells = build(job, rates)
    path = str(tmp_path_factory.mktemp("ce") / "engine.xlsx")
    wb.save(path)
    out = subprocess.run([sys.executable, os.path.join(ROOT, "tools", "recalc.py"), path, "60"],
                         capture_output=True, text=True)
    assert '"status": "success"' in out.stdout, out.stdout + out.stderr
    return load_workbook(path, data_only=True)["Quote Builder"], cells


def test_matches_reference(engine):
    ws, cells = engine
    for key, want in EXPECTED.items():
        assert ws[cells[key]].value == pytest.approx(want, rel=1e-9), key


def test_gross_margin_not_cost_plus(engine):
    ws, _ = engine
    # Repair labor: $1,200 cost at 50% margin sells at $2,400 (cost-plus would give $1,800).
    assert ws["H15"].value == pytest.approx(2400)
    assert ws["H11"].value == pytest.approx(1600 / 0.9)


def test_checks(engine):
    ws, cells = engine
    assert all(ws[c].value == "OK" for c in cells["floor_checks"])
    assert ws[cells["rate_check"]].value == "OK"


def test_override_needs_basis(tmp_path):
    job = tmp_path / "j.yaml"
    job.write_text("quote: {number: X, date: 2026-10-26, prepared_for: a, end_client: b}\n"
                   "overrides: {energy_rate: {value: 0.2}}\n")
    with pytest.raises(ValueError):
        load(str(job))

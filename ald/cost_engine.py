"""Build an ALD leak detection + repair Cost Engine workbook from a job file.

    python -m ald.cost_engine jobs/kawasaki_maryville.yaml [-o output/]

Layout and formulas follow the Kawasaki Maryville Cost Engine (ALD-2026-1013-KM):
a Rate Card sheet of inputs, each with its basis, and a Quote Builder sheet that
is formula-driven end to end. Pricing is gross margin: Sell = Cost / (1 - margin).
"""
import argparse
import datetime as dt
import os

import yaml
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

FONT = "Arial"
INPUT = Font(name=FONT, color="0000FF")
LINK = Font(name=FONT, color="008000")
BODY = Font(name=FONT)
BOLD = Font(name=FONT, bold=True)
TITLE = Font(name=FONT, bold=True, size=13)
HEAD = Font(name=FONT, bold=True, color="FFFFFF")
YELLOW = PatternFill("solid", fgColor="FFFF00")
BLACK = PatternFill("solid", fgColor="0D0D0D")
GRAY = PatternFill("solid", fgColor="F2F2F2")
RED_FILL = PatternFill("solid", fgColor="F8D7DA")
THIN = Border(bottom=Side(style="thin", color="D8D8D8"))
WRAP = Alignment(wrap_text=True, vertical="top")

USD = '$#,##0.00;($#,##0.00);"-"'
USD0 = '$#,##0;($#,##0);"-"'
PCT = '0.0%;(0.0%);"-"'
NUM0 = '#,##0;(#,##0);"-"'

# Rate card sections, in sheet order. Keys match rate_card.yaml.
SECTIONS = [
    ("PRICING ASSUMPTIONS", "pricing", ["energy_rate", "operating_hours", "kw_per_cfm", "margin_floor",
                                        "travel_margin", "ld_labor_margin", "repair_labor_margin", "rounding"]),
    ("SYSTEM / SIZING INPUTS", "sizing", ["leak_share", "guarantee_share", "cfm_per_leak", "parts_per_leak"]),
    ("SERVICE RATE CARD (ALD COST)", "service", ["tech_day", "flight", "per_diem", "hotel", "rental_car", "air_sample"]),
    ("DAYS", "days", ["survey_days", "report_days", "repair_days", "travel_days", "techs", "weekend_nights"]),
]
LABELS = {
    "energy_rate": "Energy Rate", "operating_hours": "Operating Hours", "kw_per_cfm": "kW per CFM",
    "margin_floor": "Materials / Service Margin Floor", "travel_margin": "Travel / Expense Margin",
    "ld_labor_margin": "Labor Margin (leak detection)", "repair_labor_margin": "Repair Labor Margin",
    "rounding": "Price rounding increment", "leak_share": "Expected leak share",
    "guarantee_share": "Guaranteed repair share", "cfm_per_leak": "Avg CFM per leak (planning)",
    "parts_per_leak": "Parts cost per leak", "tech_day": "Tier 1 Field Technician",
    "flight": "Flight, roundtrip per person", "per_diem": "Per Diem per person", "hotel": "Hotel",
    "rental_car": "Rental Car (Enterprise, Corp Acct AQT962X)", "air_sample": "Air sample (ISO 8573-1) cost",
    "survey_days": "Survey on-site days", "report_days": "Report days", "repair_days": "Repair on-site days",
    "travel_days": "Travel days (in/out)", "techs": "Technicians on site",
    "weekend_nights": "Weekend nights between survey and repair",
    "share_running": "Share of equipment running, normal production", "cfm_per_hp": "CFM per HP",
}
PCT_KEYS = {"margin_floor", "travel_margin", "ld_labor_margin", "repair_labor_margin",
            "leak_share", "guarantee_share", "share_running"}
USD_KEYS = {"energy_rate", "rounding", "parts_per_leak", "tech_day", "flight", "per_diem",
            "hotel", "rental_car", "air_sample"}


def load(job_path, rate_card_path=None):
    with open(rate_card_path or os.path.join(ROOT, "rate_card.yaml")) as f:
        card = yaml.safe_load(f)
    with open(job_path) as f:
        job = yaml.safe_load(f)
    flat = {}
    for section in card.values():
        for key, item in section.items():
            flat[key] = dict(item, overridden=False)
    for key, item in (job.get("overrides") or {}).items():
        if key not in flat:
            raise KeyError(f"override '{key}' is not a rate card key")
        if "basis" not in item:
            raise ValueError(f"override '{key}' needs a basis")
        flat[key].update(value=item["value"], basis=item["basis"], overridden=True)
    return job, flat


def _fmt(key):
    if key in PCT_KEYS:
        return PCT
    if key == "energy_rate":
        return "$0.0000"
    if key in USD_KEYS:
        return USD0
    return NUM0


def build(job, rates):
    wb = Workbook()
    rc = wb.active
    rc.title = "Rate Card"
    qb = wb.create_sheet("Quote Builder")
    q = job["quote"]
    R = {}  # rate key -> absolute reference on the Rate Card

    def ref(key):
        return f"'Rate Card'!{R[key]}"

    # ---------------- Rate Card ----------------
    rc.column_dimensions["A"].width = 46
    rc.column_dimensions["B"].width = 14
    rc.column_dimensions["C"].width = 10
    rc.column_dimensions["D"].width = 110
    rc["A1"] = f"ALD RATE CARD & ASSUMPTIONS · {q['end_client'].split('·')[0].strip().upper()}"
    rc["A1"].font = TITLE
    rc["A2"] = ("Blue text on yellow = input. Every assumption carries its basis in column D. "
                "Defaults come from rate_card.yaml; rows marked JOB OVERRIDE were set in the job file.")
    rc["A2"].font = BODY
    row = 4

    def put_rate(key, item):
        nonlocal row
        rc.cell(row, 1, LABELS[key]).font = BODY
        c = rc.cell(row, 2, item["value"])
        c.font, c.fill, c.number_format = INPUT, YELLOW, _fmt(key)
        rc.cell(row, 3, item.get("unit", "")).font = BODY
        basis = ("JOB OVERRIDE. " if item.get("overridden") else "") + item["basis"]
        d = rc.cell(row, 4, basis)
        d.font, d.alignment = BODY, WRAP
        R[key] = f"$B${row}"
        row += 1

    for title, _section, keys in SECTIONS:
        rc.cell(row, 1, title).font = BOLD
        row += 1
        for key in keys:
            put_rate(key, rates[key])
        row += 1

    # Compressor inventory and sizing
    comp = job.get("compressors") or {"units": []}
    rc.cell(row, 1, f"COMPRESSOR INVENTORY AND SIZING · {comp.get('source', 'source not given')}").font = BOLD
    row += 1
    first_unit = row
    for unit in comp["units"]:
        rc.cell(row, 1, unit["name"]).font = BODY
        c = rc.cell(row, 2, unit["hp"])
        c.font, c.fill, c.number_format = INPUT, YELLOW, NUM0
        rc.cell(row, 3, "HP").font = BODY
        rc.cell(row, 4, unit.get("note", f"Per {comp.get('source', 'job file')}.")).font = BODY
        row += 1
    last_unit = row - 1
    rc.cell(row, 1, "Listed HP total").font = BODY
    rc.cell(row, 2, f"=SUM(B{first_unit}:B{last_unit})" if comp["units"] else 0).number_format = NUM0
    rc.cell(row, 3, "HP")
    rc.cell(row, 4, "Calc.")
    R["listed_hp"] = f"$B${row}"
    row += 1
    put_rate("share_running", rates["share_running"])
    rc.cell(row, 1, "Running HP").font = BODY
    rc.cell(row, 2, f"={R['listed_hp']}*{R['share_running']}").number_format = NUM0
    rc.cell(row, 3, "HP")
    rc.cell(row, 4, "Calc.")
    R["running_hp"] = f"$B${row}"
    row += 1
    put_rate("cfm_per_hp", rates["cfm_per_hp"])
    rc.cell(row, 1, "Sizing flow from HP (running HP x CFM per HP)").font = BODY
    rc.cell(row, 2, f"=ROUND({R['running_hp']}*{R['cfm_per_hp']},0)").number_format = NUM0
    rc.cell(row, 3, "CFM")
    rc.cell(row, 4, "Calc.")
    R["hp_flow"] = f"$B${row}"
    row += 1
    rc.cell(row, 1, "Measured plant flow (leave blank if none)").font = BODY
    c = rc.cell(row, 2, job.get("measured_cfm"))
    c.font, c.fill, c.number_format = INPUT, YELLOW, NUM0
    rc.cell(row, 3, "CFM")
    rc.cell(row, 4, "Enter logged average flow when available. It replaces the HP-based estimate.")
    R["measured_cfm"] = f"$B${row}"
    row += 1
    rc.cell(row, 1, "Sizing flow (plant air) used by the quote").font = BOLD
    rc.cell(row, 2, f"=IF(ISNUMBER({R['measured_cfm']}),{R['measured_cfm']},{R['hp_flow']})").number_format = NUM0
    rc.cell(row, 3, "CFM")
    rc.cell(row, 4, "Calc. Measured flow when present, otherwise running HP x CFM per HP.")
    R["sizing_flow"] = f"$B${row}"
    row += 1
    rc.cell(row, 1, "Implied savings per CFM").font = BODY
    rc.cell(row, 2, f"={R['energy_rate']}*{R['operating_hours']}*{R['kw_per_cfm']}").number_format = USD
    rc.cell(row, 3, "$/CFM/yr")
    rc.cell(row, 4, "Calc: rate x hours x kW/CFM")
    row += 1
    rc.cell(row, 1, "Expected leak load in 200 HP units").font = BODY
    rc.cell(row, 2, f"={R['sizing_flow']}*{R['leak_share']}/(200*{R['cfm_per_hp']})").number_format = "0.0"
    rc.cell(row, 3, "units")
    rc.cell(row, 4, "Calc. Savings only convert to kW when units can be turned off or unloaded.")
    row += 1
    rc.cell(row, 1, "Guaranteed CFM in 200 HP units").font = BODY
    rc.cell(row, 2, f"={R['sizing_flow']}*{R['leak_share']}*{R['guarantee_share']}/(200*{R['cfm_per_hp']})").number_format = "0.0"
    rc.cell(row, 3, "units")
    rc.cell(row, 4, "Calc.")

    # ---------------- Quote Builder ----------------
    widths = [8, 20, 70, 8, 8, 14, 12, 14, 14, 14, 14, 10]
    for i, w in enumerate(widths, start=1):
        qb.column_dimensions[chr(64 + i)].width = w
    qb["A1"] = "ALD QUOTE COST ENGINE · INTERNAL, NOT THE CUSTOMER COPY"
    qb["A1"].font = TITLE
    date = q["date"] if isinstance(q["date"], dt.date) else dt.date.fromisoformat(str(q["date"]))
    for r, (k1, v1, k2, v2) in enumerate([
        ("Quote No.", q["number"], "Date", date),
        ("Prepared For", q["prepared_for"], "Valid Through", f"=G3+{int(q.get('valid_days', 30))}"),
        ("End Client", q["end_client"], "Prepared By", q.get("prepared_by", "Dominick Ruggeri")),
    ], start=3):
        qb.cell(r, 1, k1).font = BOLD
        qb.cell(r, 3, v1).font = INPUT
        qb.cell(r, 6, k2).font = BOLD
        c = qb.cell(r, 7, v2)
        c.font = INPUT if r != 4 else BODY
        if r in (3, 4):
            c.number_format = "m/d/yy"

    hdr = ["Option", "Category", "Description", "Qty", "Unit", "Unit Cost ($)", "Target Margin %",
           "Unit Sell ($)", "Ext. Cost ($)", "Ext. Sell ($)", "Profit ($)", "Margin %"]
    for i, h in enumerate(hdr, start=1):
        c = qb.cell(8, i, h)
        c.font, c.fill, c.alignment = HEAD, BLACK, Alignment(wrap_text=True, horizontal="center")

    route = q.get("travel_route", "Livermore to site")
    T, S, RP, TR, WK = (ref(k) for k in ("techs", "survey_days", "repair_days", "travel_days", "weekend_nights"))
    # Leak load CFM lives in the Savings block; parts qty references it (cell fixed below).
    lines = [
        (1, "Leak Detection Labor", "Tier 1 techs: (techs x survey on-site days) + report days",
         f"={T}*{S}+{ref('report_days')}", "day", "tech_day", "ld_labor_margin"),
        (1, "Air Sample", "Complimentary ISO 8573-1 air sample (cost absorbed, $0 sell)",
         1, "ea", "air_sample", None),
        (1, "Travel", f"Flight RT, per tech ({route})", f"={T}", "ea", "flight", "travel_margin"),
        (1, "Travel", "Hotel (techs x survey on-site days, arrive evening before)",
         f"={T}*{S}", "night", "hotel", "travel_margin"),
        (1, "Travel", "Per diem (techs x (on-site days + fly-in day))",
         f"={T}*({S}+{TR})", "day", "per_diem", "travel_margin"),
        (1, "Travel", "Rental car (Enterprise, one week, shared car)", 1, "wk", "rental_car", "travel_margin"),
        (2, "Repair Labor", "Tier 1 techs: techs x repair on-site days",
         f"={T}*{RP}", "day", "tech_day", "repair_labor_margin"),
        (2, "Parts", "Repair parts: est. leaks in full expected leak load (leak load CFM / avg CFM per leak)",
         "PARTS", "leak", "parts_per_leak", "margin_floor"),
        (2, "Travel", "Added hotel nights (techs x (repair days + weekend nights), same trip)",
         f"={T}*({RP}+{WK})", "night", "hotel", "travel_margin"),
        (2, "Travel", "Added per diem (techs x (repair days + weekend days))",
         f"={T}*({RP}+{WK})", "day", "per_diem", "travel_margin"),
        (2, "Travel", "Added rental car (repair weeks beyond the survey week, same car)",
         f"=ROUNDUP({RP}/5,0)", "wk", "rental_car", "travel_margin"),
    ]
    for extra in job.get("extra_lines") or []:
        lines.append((extra["option"], extra["category"], extra["description"], extra["qty"],
                      extra["unit"], extra["cost"], extra["margin"]))

    first = 9
    last = first + len(lines) - 1
    tot = last + 3          # OPTION TOTALS header row
    survey_row, repair_row, both_row = tot + 2, tot + 3, tot + 4
    sv = both_row + 3       # first SAVINGS row
    LEAK_CFM = f"$C${sv + 1}"

    for i, (opt, cat, desc, qty, unit, cost_key, margin_key) in enumerate(lines):
        r = first + i
        qb.cell(r, 1, opt)
        qb.cell(r, 2, cat)
        qb.cell(r, 3, desc).alignment = Alignment(wrap_text=True)
        if qty == "PARTS":
            qty = f"=ROUNDUP({LEAK_CFM}/{ref('cfm_per_leak')},0)"
        qb.cell(r, 4, qty).number_format = NUM0
        qb.cell(r, 5, unit)
        if isinstance(cost_key, str) and cost_key in R:
            qb.cell(r, 6, f"={ref(cost_key)}").font = LINK
        else:
            qb.cell(r, 6, cost_key).font = INPUT
        qb.cell(r, 6).number_format = USD
        if margin_key is None:      # complimentary line: sells at $0
            qb.cell(r, 7, 0).font = INPUT
            qb.cell(r, 8, 0).font = INPUT
            qb.cell(r, 8).comment = Comment("Complimentary. Sell fixed at $0, cost absorbed.", "ALD")
        else:
            if isinstance(margin_key, str):
                qb.cell(r, 7, f"={ref(margin_key)}").font = LINK
            else:
                qb.cell(r, 7, margin_key).font = INPUT
            qb.cell(r, 8, f"=IFERROR(F{r}/(1-G{r}),0)")
        qb.cell(r, 7).number_format = PCT
        qb.cell(r, 8).number_format = USD
        qb.cell(r, 9, f"=D{r}*F{r}").number_format = USD
        qb.cell(r, 10, f"=D{r}*H{r}").number_format = USD
        qb.cell(r, 11, f"=J{r}-I{r}").number_format = USD
        qb.cell(r, 12, f"=IFERROR(K{r}/J{r},0)").number_format = PCT
        for col in range(1, 13):
            qb.cell(r, col).border = THIN
            if qb.cell(r, col).font == Font():
                qb.cell(r, col).font = BODY

    qb.cell(tot, 1, "OPTION TOTALS").font = BOLD
    for i, h in enumerate(["Option", "Ext. Cost ($)", "Ext. Sell (calc)", "CUSTOMER PRICE (rounded up)",
                           "Profit ($)", "Margin % (on customer price)", "Floor check (>=floor)"], start=1):
        c = qb.cell(tot + 1, i, h)
        c.font, c.fill, c.alignment = HEAD, BLACK, Alignment(wrap_text=True)
    rng = f"$A${first}:$A${last}"
    for r, label, opt in ((survey_row, "Survey", 1), (repair_row, "Repair add-on", 2)):
        qb.cell(r, 1, label)
        qb.cell(r, 2, f"=SUMIF({rng},{opt},$I${first}:$I${last})")
        qb.cell(r, 3, f"=SUMIF({rng},{opt},$J${first}:$J${last})")
        qb.cell(r, 4, f"=CEILING(C{r},{ref('rounding')})")
    qb.cell(both_row, 1, "Survey + repair add-on")
    for col in "BCD":
        qb[f"{col}{both_row}"] = f"={col}{survey_row}+{col}{repair_row}"
    for r in (survey_row, repair_row, both_row):
        qb.cell(r, 5, f"=D{r}-B{r}")
        qb.cell(r, 6, f"=IFERROR(E{r}/D{r},0)").number_format = PCT
        qb.cell(r, 7, f'=IF(F{r}>={ref("margin_floor")},"OK","BELOW FLOOR")')
        for col in range(2, 6):
            qb.cell(r, col).number_format = USD
        qb.cell(r, 4).font = BOLD
    qb.conditional_formatting.add(f"G{survey_row}:G{both_row}",
                                  FormulaRule(formula=[f'G{survey_row}="BELOW FLOOR"'], fill=RED_FILL))

    # Savings & ROI
    E, H, K = ref("energy_rate"), ref("operating_hours"), ref("kw_per_cfm")
    D1, D2, D3 = f"$D${survey_row}", f"$D${repair_row}", f"$D${both_row}"
    savings = [
        ("Sizing flow (CFM)", f"={ref('sizing_flow')}", NUM0, "Rate Card"),
        ("Expected leak load (CFM)", f"=ROUND(C{sv}*{ref('leak_share')},0)", NUM0, "Flow x leak share"),
        ("Leak load kWh/yr", f"=C{sv+1}*{K}*{H}", NUM0, "CFM x kW/CFM x hrs"),
        ("Potential annual savings, full leak load ($)", f"=C{sv+2}*{E}", USD0, "kWh x rate"),
        ("Guaranteed CFM (repair add-on)", f"=ROUND(C{sv+1}*{ref('guarantee_share')},0)", NUM0,
         "Leak load x guarantee share. ONE driver for the repair add-on."),
        ("Guaranteed kW", f"=C{sv+4}*{K}", "0.0", ""),
        ("Guaranteed bhp", f"=C{sv+5}/0.746", NUM0, "kW / 0.746"),
        ("Guaranteed kWh/yr", f"=C{sv+5}*{H}", NUM0, ""),
        ("Guaranteed annual savings ($)", f"=C{sv+7}*{E}", USD0, ""),
        ("Cost of delay, full leak load ($/month)", f"=C{sv+3}/12", USD0, ""),
        ("Survey payback on potential savings (months)", f"=IFERROR({D1}/C{sv+3}*12,0)", "0.0", ""),
        ("Repair add-on payback on guaranteed savings (months)", f"=IFERROR({D2}/C{sv+8}*12,0)", "0.0", ""),
        ("Survey + repair add-on payback on guaranteed savings (months)", f"=IFERROR({D3}/C{sv+8}*12,0)", "0.0",
         "Conservative: no credit for leaks found but not repaired"),
        ("Leak load bhp", f"=C{sv+1}*{K}/0.746", NUM0, ""),
        ("Leak load kW", f"=C{sv+1}*{K}", NUM0, ""),
    ]
    qb.cell(sv - 1, 1, "SAVINGS & ROI").font = BOLD
    for i, (label, formula, fmt, note) in enumerate(savings):
        r = sv + i
        qb.cell(r, 1, label)
        qb.cell(r, 3, formula).number_format = fmt
        qb.cell(r, 5, note)

    mh = sv + len(savings) + 1
    qb.cell(mh, 1, "MARGIN HEADROOM").font = BOLD
    headroom = [
        ("Price at the margin floor, survey + repair add-on",
         f"=CEILING($B${both_row}/(1-{ref('margin_floor')}),{ref('rounding')})", USD0, ""),
        ("Customer price above the floor", f"={D3}-C{mh+1}", USD0, ""),
        ("Guarantee if the survey measures only half the expected leak load",
         f"=ROUND(C{sv+1}*0.5*{ref('guarantee_share')},0)", NUM0,
         "Guarantee falls to the guarantee share of the measured leak load."),
        ("Energy rate check", f'=IF({E}=0.13,"PLACEHOLDER RATE: set the plant rate in the job file","OK")',
         None, "Flags the rate card default left in place."),
    ]
    for i, (label, formula, fmt, note) in enumerate(headroom, start=1):
        r = mh + i
        qb.cell(r, 1, label)
        c = qb.cell(r, 3, formula)
        if fmt:
            c.number_format = fmt
        qb.cell(r, 5, note)
    rate_chk = f"C{mh+4}"
    qb.conditional_formatting.add(rate_chk, FormulaRule(formula=[f'{rate_chk}<>"OK"'], fill=RED_FILL))

    se = mh + 6
    qb.cell(se, 1, "SENSITIVITY (internal only, not on the customer copy)").font = BOLD
    qb.cell(se + 1, 1, "Alt hours per year (input)")
    qb.cell(se + 1, 2, 6000).font = INPUT
    qb.cell(se + 1, 2).fill = YELLOW
    qb.cell(se + 2, 1, "Alt demand as share of nameplate sizing (input)")
    qb.cell(se + 2, 2, 0.7).font = INPUT
    qb.cell(se + 2, 2).fill = YELLOW
    qb.cell(se + 2, 2).number_format = PCT
    sens = [
        ("Full leak load savings at alt hours ($)", f"=C{sv+1}*{K}*B{se+1}*{E}"),
        ("Guaranteed savings at alt hours ($)", f"=C{sv+4}*{K}*B{se+1}*{E}"),
        ("Full leak load savings at alt demand share ($)", f"=C{sv+3}*B{se+2}"),
        ("Guaranteed savings at alt demand share ($)", f"=C{sv+8}*B{se+2}"),
    ]
    for i, (label, formula) in enumerate(sens, start=3):
        qb.cell(se + i, 1, label)
        qb.cell(se + i, 3, formula).number_format = USD0

    roi = se + 8
    qb.cell(roi, 1, "FIRST YEAR ROI (net first year savings / price)").font = BOLD
    for i, (label, sav, price) in enumerate([
        ("ROI, survey on full leak load (all leaks repaired)", f"C{sv+3}", D1),
        ("ROI, repair add-on on guaranteed savings", f"C{sv+8}", D2),
        ("ROI, survey + repair on full leak load", f"C{sv+3}", D3),
        ("ROI, survey + repair on guaranteed savings", f"C{sv+8}", D3),
    ], start=1):
        qb.cell(roi + i, 1, label)
        qb.cell(roi + i, 3, f"=IFERROR(({sav}-{price})/{price},0)").number_format = "0%"
        qb.cell(roi + i, 5, "(Savings - price) / price, one year")

    for ws in (rc, qb):
        for row_cells in ws.iter_rows():
            for c in row_cells:
                if c.value is not None and c.font.name != FONT:
                    c.font = Font(name=FONT, bold=c.font.bold, color=c.font.color)
    qb.freeze_panes = "A9"
    wb.active = 1
    cells = dict(survey=f"D{survey_row}", repair=f"D{repair_row}", total=f"D{both_row}",
                 leak_cfm=f"C{sv+1}", guaranteed_cfm=f"C{sv+4}", potential=f"C{sv+3}",
                 guaranteed_savings=f"C{sv+8}", rate_check=rate_chk,
                 floor_checks=[f"G{survey_row}", f"G{repair_row}", f"G{both_row}"])
    return wb, cells


def main():
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("job")
    ap.add_argument("-o", "--outdir", default=os.path.join(ROOT, "output"))
    ap.add_argument("--rate-card")
    a = ap.parse_args()
    job, rates = load(a.job, a.rate_card)
    wb, _ = build(job, rates)
    os.makedirs(a.outdir, exist_ok=True)
    q = job["quote"]
    name = f"ALD Quote - {q['number']} - Cost Engine.xlsx"
    path = os.path.join(a.outdir, name)
    wb.save(path)
    print(path)


if __name__ == "__main__":
    main()
